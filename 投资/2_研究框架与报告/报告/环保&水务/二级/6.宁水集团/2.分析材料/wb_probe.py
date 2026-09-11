import openpyxl
p="/Users/zhangarmanda/PycharmProjects/Personal/投资/2_研究框架与报告/报告/环保&水务/二级/6.宁水集团/2.分析材料/2_2026-08_宁水集团_取数底稿.xlsx"
wb=openpyxl.load_workbook(p)
print("SHEETS:", wb.sheetnames)
for ws in wb:
    print("\n===== %s  rows=%d cols=%d" % (ws.title, ws.max_row, ws.max_column))
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,60)):
        cells=[]
        for c in r:
            if c.value is not None:
                fill = c.fill.start_color.rgb if c.fill and c.fill.start_color else None
                tag = ""
                if fill and isinstance(fill,str):
                    if fill.endswith("FFF2CC"): tag="[Y]"
                    elif fill.endswith("E7E6E6"): tag="[G]"
                    elif fill.endswith("FCE4D6"): tag="[O]"
                cells.append("%s%s=%s" % (c.coordinate, tag, str(c.value)[:60]))
        if cells: print(" | ".join(cells))
