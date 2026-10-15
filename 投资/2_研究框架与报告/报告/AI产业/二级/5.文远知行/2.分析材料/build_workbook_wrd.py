"""从已核验披露构建公式底稿；限定公式校验不冒充 Office 原生重算。"""
from pathlib import Path
import ast
import copy
import csv
import json
import math
import operator
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties

HERE = Path(__file__).resolve().parent
ROOT = next(p for p in HERE.parents if (p / '投资').is_dir())
OUT = HERE / '2_2026-09_文远知行_取数底稿.xlsx'
wb = load_workbook(ROOT / '投资/2_研究框架与报告/模板/取数底稿模板_二级市场.xlsx')
for ws in wb:
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in ws:
        for cell in row:
            cell.value = None
            cell.comment = None
raw = wb.create_sheet('7_原始输入')
raw.append(['键', '原值', '单位', '来源及页码', '证据等级/口径'])
log = wb['6_取数日志']
log.delete_rows(1, log.max_row)
for i, text in enumerate(['采集日', '指标/输入单元格', '原值', '来源及页码', '口径'], 1):
    log.cell(1, i, text)
manifest = json.loads((HERE / 'source_manifest.json').read_text())
market = json.loads((HERE / 'market_audit.json').read_text())
refs = {}

def inp(key, value, source, unit='人民币千元', note='披露'):
    if key in refs:
        raise ValueError(key)
    raw.append([key, value, unit, source, note])
    r = raw.max_row
    refs[key] = f"'7_原始输入'!B{r}"
    raw.cell(r, 2).comment = Comment(source + '\n' + note, '来源')
    log.append(['2026-09-17', key + f' / 7_原始输入!B{r}', value, source, note + '；单位：' + unit])
    return refs[key]

def ref(key):
    return refs[key]

def put(ws, r, c, value):
    ws.cell(r, c, value)

def source(name, pages):
    return f'{name} PDF页{pages}；' + manifest.get(name + '.pdf', {}).get('url', '')

US = source('2024_US_424B4', '279/F-3；286–287/F-10–11')
Y24 = source('FY2024_20F', '169/F-3；176–177/F-10–11')
Y25 = source('FY2025_20F', '171/F-4；178/F-11；111/108')
H26 = source('2026H1_interim', '47/46；53–54/52–53')
HK = source('2025_HK_prospectus', '592–593/I-16–17')
annual = {
    '收入': [138172,527543,401844,361134,684587],
    '产品': [101597,337717,54190,87710,359843],
    '服务': [36575,189826,347654,273424,324744],
    '毛利': [51660,232545,183476,110715,206822],
    '研发': [-443178,-758565,-1058395,-1091357,-1372191],
    '行政': [-107119,-237236,-625369,-1138802,-596060],
    '销售': [-12225,-23574,-41447,-53566,-73628],
    '经营亏损': [-500496,-779230,-1566202,-2185183,-1846770],
    '税前亏损': [-1007251,-1298496,-1946235,-2510940,-1650800],
    '归母净亏损': [-1007251,-1298496,-1949101,-2516808,-1654900],
    '股权支付': [55959,325429,931784,1187867,449983],
    '调整后亏损': [-426757,-401683,-501680,-801853,-1246739],
    '融资成本': [-6917,-4202,-3490,-3451,-9128],
    'OCF': [-506667,-670381,-474890,-593595,-1321734],
    '固定资产购置': [25156,80812,36650,84004,247598],
    '无形资产购置': [481,1881,304,1504,848],
    '投资现金流': [460903,-2202414,-546944,325505,1505881],
    '融资现金流': [2603053,2184588,446954,2823875,2296444],
    '现金': [2725568,2233691,1661152,4268300,6666304],
}
half = {
    '收入': [182869,150298,199615,345857],
    '产品': [18553,21045,69281,112783],
    '服务': [164316,129253,130334,233074],
    '毛利': [83975,54789,61059,126473],
    '研发': [-376121,-517210,-644635,-797654],
    '行政': [-217101,-208293,-278942,-152066],
    '销售': [-14619,-22784,-27780,-51891],
    '经营亏损': [-538270,-698983,-890077,-852910],
    '归母净亏损': [-723092,-881700,-791517,-789750],
    '融资成本': [-1784,-1356,-3292,-5089],
    'OCF': [-223742,-327180,-663396,-715677],
    '固定资产购置': [15906,33272,134354,235442],
    '无形资产购置': [127,0,117,808],
    '投资现金流': [-555576,453236,218699,-960615],
    '融资现金流': [193356,-8877,3899,-544316],
    '现金': [1659970,1828943,3836137,4275014],
}
for metric, values in annual.items():
    for year, value in zip(range(2021,2026), values):
        s = US if year == 2021 else Y24 if year == 2022 else Y25
        if metric == '调整后亏损' and year == 2021:
            s = source('2024_US_424B4', '40–43/摘要Non-IFRS调节')
        inp(f'FY{year}_{metric}',value,s,note='审计披露' if metric != '调整后亏损' else '管理层Non-IFRS；非中国扣非口径')
for metric, values in half.items():
    for year, value in zip(range(2023,2027), values):
        s = source('2024_US_424B4','355/F-79；361–362/F-85–86') if year <= 2024 else H26
        if year == 2024 and metric in ['OCF','融资现金流']:
            s = HK
        inp(f'H1{year}_{metric}',value,s,note='半年披露；2024 OCF采用港股招股书后列比较数，较美招股书差378千元')

bs = {
    '总资产':[None,5385391,5614258,7693770,8935925,7755416],
    '归母权益':[None,-2082116,-3051918,7066019,7900174,6395302],
    '现金':[2725568,2233691,1661152,4268300,6666304,4275014],
    '短期贷款':[None,0,0,30019,324263,485047],
    '长期贷款':[None,0,0,50040,0,0],
    '租赁流动':[None,32009,31098,36900,31920,35390],
    '租赁非流动':[None,35864,22309,26059,23241,95397],
    '贸易应收净额':[None,236390,266933,252607,462135,455765],
    '合同资产净额':[None,92597,82826,28005,23305,23964],
    '定期存款':[None,1057292,2550279,620148,301401,1099805],
    '流动FVTPL':[None,1218524,317042,1685146,144252,2293],
    '非流动FVTPL':[None,None,0,56919,188083,252411],
    '受限现金流动':[None,1393,10194,4814,12910,10420],
    '受限现金非流动':[None,11004,1575,9669,6487,10988],
    '存货':[None,156005,218220,204705,321021,439112],
    '总负债':[None,7467507,8666176,627751,1035751,1360114],
    '流动资产':[None,5073473,5370023,7287995,8210324,6628076],
    '流动负债':[None,None,8591413,542489,1001301,1236732],
    '应收拨备':[None,None,None,None,69795,66058],
    '合同资产拨备':[None,None,None,None,23910,23495],
    '固定资产':[None,113878,98574,178179,378769,573233],
    '回售期权负债':[None,39812,40449,41099,43300,31681],
}
periods = ['FY2021','FY2022','FY2023','FY2024','FY2025','H12026']
for metric, values in bs.items():
    for period, value in zip(periods, values):
        if value is None or f'{period}_{metric}' in refs:
            continue
        s = source('2024_US_424B4','281–282/F-5–6') if period == 'FY2022' else source('FY2024_20F','171–172/F-5–6') if period == 'FY2023' else source('FY2025_20F','173–174/F-6–7')
        if period == 'H12026' or metric in ['应收拨备','合同资产拨备']:
            s = source('2026H1_interim','49–50/48–49；66–67/65–66')
        inp(f'{period}_{metric}',value,s)
for key,val in {'股份A':972508041,'股份B':54814423,'库存股':40454099,'ADS比例':3,'期权':90902378,'RSU2018':19094091,'RSU2026':7686812,'预发行激励股份':75000000}.items():
    inp(key,val,source('2026H1_interim','43/42；73/72；75–77/74–76'),unit='股；ADS比例为股/ADS')
for key,val in {'半年回购':684809,'半年租赁本金':24578,'半年租赁利息':1330,'半年贷款利息':3711,'半年汇率影响':-170682,'FY2025租赁本金':46629,'FY2025租赁利息':2745,'FY2025贷款利息':4063,'H12025租赁本金':26810,'H12025租赁利息':1660,'H12025贷款利息':1280}.items():
    inp(key,val,H26 if not key.startswith('FY') else Y25)
for key,val in {'汇率':6.7851,'要求回报':0.10,'期限年':10,'半年月数':6,'年内月数':12,'交易日年基数':244,'研究观察价':3.30,'压力烧钱倍率':1.5}.items():
    inp(key,val,'PONY_2026Q2公告汇率说明（2026-06-30美联储H.10）' if key=='汇率' else 'Noah v1.2 [C1]/[C3]；其余为本报告显式分析假设或单位常数',unit='参数',note='披露汇率' if key=='汇率' else '假设/单位常数；非公司指引')
for ticker,m in market.items():
    for key,value in [('价格',m['close']),('市值',m['snapshot']['f116']),('股数',m['snapshot']['f84']),('PB接口',m['snapshot']['f167'])]:
        inp(f'{ticker}_{key}',value,m['snapshot_url']+'；2026-09-16收盘',unit='USD、股/ADS；PB接口百分缩放',note='行情；非完全摊薄股本')
pony = {'收入':70470,'去年收入':35434,'毛利':11911,'净亏损':-98861,'归母亏损':-110251,'NCI利润':11390,'OCF':-118215,'去年OCF':-79570,'capex':44618,'去年capex':14464,'现金':327111,'短期投资':787008,'受限现金流动':4678,'受限现金非流动':117,'长期投资':375660,'归母权益':1575534,'总负债':93005,'Robotaxi':20643,'Robotruck':23524,'Intelligent':26303}
for key,value in pony.items():
    inp('PONY_'+key,value,'PONY_2026Q2_webfetch.txt；官方公告三表及Non-GAAP调节表；'+manifest['PONY_2026Q2_webfetch.txt']['url'],unit='USD千元',note='半年未经审计；capex精确值已核，Q1 12455+Q2 32163=44618')
segments = {'Robotaxi及相关服务':[29379,47832,147993],'Robobus':[94914,79688,231473],'Robosweeper':[7642,55320,96083],'Robovan':[7540,13393,10493],'其他技术服务':[262369,164901,198545],'海外':[23832,49364,199761]}
for name,values in segments.items():
    for y,v in zip(range(2023,2026),values):
        inp(f'{y}_{name}',v,source('FY2025_20F','196–197/F-29–30；附注4/5'))
for name,vals in {'L4':[123932,179475],'其他业务':[75683,166382],'海外':[46929,119299]}.items():
    for y,v in zip([2025,2026],vals):
        inp(f'H1{y}_{name}',v,source('2026H1_interim','57–59/56–58；附注3/4'))

ws=wb['0_说明']
notes=['文远知行｜2026-09研究底稿｜用户指定版本 Sonus','金额展示为人民币百万，同行表USD百万；原始输入保留千元。','蓝字=原始输入，黑字=同表公式，绿字=跨表公式；黄色=假设/待核。','期间：FY2021–FY2025及H1 2023–2026；仅两份独立20-F，以招股书补历史。','FY2021完整集团资产负债表未取得，空缺标“未取得”，不将VIE资产负债表替代集团。','文远全期IFRS；PONY为US GAAP。车辆级UE不等于公司级FCF。','最新三表：2026中报PDF47、49–50、53–54；审阅而非审计。','FCF=OCF−购置固定资产−购置无形资产；另列扣租赁本金/利息及贷款利息的压力口径。','历史半年OCF采用较新HK招股书：H1 2024为−327.180，旧美招股书为−327.558；差额仍待附注解释。','归母损益不等于IPO前普通股股东损益：FY2023视同分派32.767未扣在归母行。','所有输入均在7_原始输入及6_取数日志；所有派生数均为公式。','三锚点：FY2024收入/归母/总资产与SEC ifrs-full逐值一致；SEC缓存未覆盖最新年度，不宣称最新也已双源核。','行情统一2026-09-16收盘；价格分位只作上市以来位置，不是估值分位。','现金池按账面值：不受限现金+定存+流动金融资产，不代表全部等同无风险现金。','跑道=期末现金池÷历史月均现金消耗，为静态压力刻度，不是半年度×2全年预测。','PS情景是2036年收入/权益PS反推；零终值净现金加成；稀释独立显式输入，不给收益保证。','本机没有LibreOffice；限定算术公式逐格求值并写缓存、开启Excel打开时重算；未获Office原生重算验证。','原始来源、SHA256见source_manifest.json；公式测试及勾稽见workbook_validation.json。']
for i,n in enumerate(notes,1): put(ws,i,1,n)
ws.column_dimensions['A'].width=135

ws=wb['1_分部拆解']
for c,t in enumerate(['收入拆解（人民币百万）']+periods+['口径'],1):put(ws,1,c,t)
for r,m in enumerate(['产品','服务','收入'],2):
    put(ws,r,1,m)
    for c,p in enumerate(periods,2): put(ws,r,c,'='+ref(f'{p}_{m}')+'/1000')
for c in range(2,8):
    x=col(c)
    for r,expr in [(6,f'{x}2/{x}4'),(7,f'{x}3/{x}4'),(9,f'{x}2+{x}3-{x}4')]:put(ws,r,c,'='+expr)
put(ws,6,1,'产品占比');put(ws,7,1,'服务占比');put(ws,9,1,'收入勾稽差')
put(ws,11,1,'业务线而非独立经营分部；Robotaxi含车辆/技术支持，不是纯车费')
for c,t in enumerate(['车型业务（人民币百万）','FY2023','FY2024','FY2025','FY2025占比'],1):put(ws,12,c,t)
for r,m in enumerate(list(segments)[:5],13):
    put(ws,r,1,m)
    for c,y in enumerate(range(2023,2026),2):put(ws,r,c,'='+ref(f'{y}_{m}')+'/1000')
    put(ws,r,5,f'=D{r}/F4')
for c in range(2,5):put(ws,18,c,f'=SUM({col(c)}13:{col(c)}17)')
put(ws,18,1,'业务线合计')
for c,t in enumerate(['最新分类（不可硬接车型）','H12025','H12026','同比','H12026占比'],1):put(ws,21,c,t)
for r,m in enumerate(['L4','其他业务','海外'],22):
    put(ws,r,1,m)
    for c,y in enumerate([2025,2026],2):put(ws,r,c,'='+ref(f'H1{y}_{m}')+'/1000')
    put(ws,r,4,f'=C{r}/B{r}-1');put(ws,r,5,f'=C{r}/G4')
put(ws,26,1,'分业务毛利率/纯车费/单车完整成本：未披露，不能用集团毛利率分摊')

ws=wb['2_利润表五年']
metrics=['收入','毛利','销售行政','融资成本','研发','税前亏损','归母净亏损','经营亏损','股权支付','调整后亏损','毛利率','研发/收入','营业亏损率','收入同比']
for c,t in enumerate(['利润表（人民币百万）']+periods+['TTM至2026-06'],1):put(ws,1,c,t)
for r,m in enumerate(metrics,2):
    put(ws,r,1,m)
    for c,p in enumerate(periods,2):
        if m=='销售行政':f='=('+ref(f'{p}_销售')+'+'+ref(f'{p}_行政')+')/1000'
        elif m in ['毛利率','研发/收入','营业亏损率','收入同比']:
            x=col(c)
            f={'毛利率':f'={x}3/{x}2','研发/收入':f'=-{x}6/{x}2','营业亏损率':f'={x}9/{x}2','收入同比':f'={x}2/{col(c-1)}2-1'}[m]
            if m=='收入同比' and c==2:f='未取得FY2020'
            if m=='收入同比' and c==7:f='='+ref('H12026_收入')+'/'+ref('H12025_收入')+'-1'
        else:f='='+ref(f'{p}_{m}')+'/1000' if f'{p}_{m}' in refs else '未录入；不估填'
        put(ws,r,c,f)
    if m in annual and m in half:put(ws,r,8,'=('+ref('FY2025_'+m)+'-'+ref('H12025_'+m)+'+'+ref('H12026_'+m)+')/1000')
put(ws,12,8,'=H3/H2')
put(ws,18,1,'调整后亏损为公司Non-IFRS，不等于A股扣非；FY2025净亏损改善但调整后亏损扩大。')

ws=wb['3_现金流照妖镜']
for c,t in enumerate(['现金流（人民币百万）']+periods+['五年合计','TTM至2026-06'],1):put(ws,1,c,t)
cfmetrics=['OCF','固定资产购置','无形资产购置','资本开支合计','FCF','投资现金流','融资现金流','现金']
for r,m in enumerate(cfmetrics,2):
    put(ws,r,1,m)
    for c,p in enumerate(periods,2):
        x=col(c)
        f=f'={x}3+{x}4' if m=='资本开支合计' else f'={x}2-{x}5' if m=='FCF' else '='+ref(f'{p}_{m}')+'/1000'
        put(ws,r,c,f)
    put(ws,r,8,f'=SUM(B{r}:F{r})' if m!='现金' else '时点值不可相加')
    put(ws,r,9,f'=F{r}-C{r+12}+G{r}' if m in ['资本开支合计','FCF'] else '=('+ref('FY2025_'+m)+'-'+ref('H12025_'+m)+'+'+ref('H12026_'+m)+')/1000' if m!='现金' else '=G9')
for c,t in enumerate(['历年半年（人民币百万）','H12023','H12024','H12025','H12026'],1):put(ws,13,c,t)
for r,m in enumerate(['OCF','固定资产购置','无形资产购置','资本开支合计','FCF','投资现金流','融资现金流','现金'],14):
    put(ws,r,1,m)
    for c,p in enumerate(['H12023','H12024','H12025','H12026'],2):
        x=col(c)
        f=f'={x}15+{x}16' if m=='资本开支合计' else f'={x}14-{x}17' if m=='FCF' else '='+ref(f'{p}_{m}')+'/1000'
        put(ws,r,c,f)
put(ws,5,9,'=F5-D17+G5');put(ws,6,9,'=F6-D18+G6')
put(ws,24,1,'H1 2026回购（不进FCF）');put(ws,24,2,'='+ref('半年回购')+'/1000')
put(ws,25,1,'H1 2026租赁本金/利息+贷款利息');put(ws,25,2,'=('+ref('半年租赁本金')+'+'+ref('半年租赁利息')+'+'+ref('半年贷款利息')+')/1000')
put(ws,26,1,'H1 2026全现金支出压力（不含融资/资金调配）');put(ws,26,2,'=-G6+B24+B25')
put(ws,27,1,'现金桥勾稽差');put(ws,27,2,'=F9+G2+G7+G8+'+ref('半年汇率影响')+'/1000-G9')
put(ws,29,1,'亏损/亏损的OCF净利比不作利润质量判断；H1 2024后列OCF与融资净额分别变378，净现金变动不变。')

ws=wb['4_资产负债偿债']
for c,t in enumerate(['资产负债（人民币百万）']+periods+['口径'],1):put(ws,1,c,t)
for r,(m,values) in enumerate(bs.items(),2):
    put(ws,r,1,m)
    for c,(p,v) in enumerate(zip(periods,values),2):put(ws,r,c,'='+ref(f'{p}_{m}')+'/1000' if v is not None else '未取得')
for r,name in [(26,'资产负债表勾稽差'),(27,'流动性池：现金+定存+流动FVTPL'),(28,'有息债务含租赁'),(29,'现金/短贷'),(30,'贸易应收+合同资产/资产'),(31,'贸易应收拨备率'),(32,'合同资产拨备率'),(33,'两类合计拨备率')]:put(ws,r,1,name)
for c in range(3,8):
    x=col(c)
    put(ws,26,c,f'={x}2-{x}3-{x}17');put(ws,27,c,f'={x}4+{x}11+{x}12')
    put(ws,28,c,f'=SUM({x}5:{x}8)')
    if c>=5:put(ws,29,c,f'={x}4/{x}5')
    put(ws,30,c,f'=({x}9+{x}10)/{x}2')
    if c>=6:
        put(ws,31,c,f'={x}20/({x}9+{x}20)');put(ws,32,c,f'={x}21/({x}10+{x}21)')
        put(ws,33,c,f'=({x}20+{x}21)/({x}9+{x}10+{x}20+{x}21)')

ws=wb['5_估值对比']
heads=['同日价、H1 2026；USD百万','收盘USD','市值','H1收入','H1毛利率','H1 OCF','H1 FCF','现金等价物','流动性池','静态FCF月数','归母PB']
for c,t in enumerate(heads,1):put(ws,1,c,t)
for r,ticker in enumerate(['WRD','PONY','AUR'],2):
    put(ws,r,1,ticker);put(ws,r,2,'='+ref(ticker+'_价格'));put(ws,r,3,'='+ref(ticker+'_市值')+'/1000000')
put(ws,2,4,'='+ref('H12026_收入')+'/1000/'+ref('汇率'))
put(ws,2,5,"='2_利润表五年'!G12")
put(ws,2,6,"='3_现金流照妖镜'!G2/"+ref('汇率'))
put(ws,2,7,"='3_现金流照妖镜'!G6/"+ref('汇率'))
put(ws,2,8,"='4_资产负债偿债'!G4/"+ref('汇率'))
put(ws,2,9,"='4_资产负债偿债'!G27/"+ref('汇率'))
put(ws,2,10,'=I2/(-G2/'+ref('半年月数')+')')
put(ws,2,11,"=C2/('4_资产负债偿债'!G3/"+ref('汇率')+')')
for c,k in [(4,'收入'),(6,'OCF'),(8,'现金')]:put(ws,3,c,'='+ref('PONY_'+k)+'/1000')
put(ws,3,5,'='+ref('PONY_毛利')+'/'+ref('PONY_收入'))
put(ws,3,7,'=('+ref('PONY_OCF')+'-'+ref('PONY_capex')+')/1000')
put(ws,3,9,'=('+ref('PONY_现金')+'+'+ref('PONY_短期投资')+')/1000')
put(ws,3,10,'=I3/(-G3/'+ref('半年月数')+')')
put(ws,3,11,'=C3/('+ref('PONY_归母权益')+'/1000)')
for c in range(4,12):put(ws,4,c,'未取财务；不可排序')
put(ws,6,1,'WRD为IFRS，另加融资中租赁/利息后可支配FCF压力（USD百万）')
put(ws,6,2,"=('3_现金流照妖镜'!G6-'3_现金流照妖镜'!B25)/"+ref('汇率'))
put(ws,7,1,'WRD上述支出压力跑道（月）');put(ws,7,2,'=I2/(-B6/'+ref('半年月数')+')')
put(ws,8,1,'WRD现金等价物跑道（月）');put(ws,8,2,'=H2/(-G2/'+ref('半年月数')+')')
put(ws,9,1,'PONY现金等价物跑道（月）');put(ws,9,2,'=H3/(-G3/'+ref('半年月数')+')')
put(ws,11,1,'流动性池不含受限、长期投资；短期投资与定存信用/价格风险不同，不能称完全同风险现金。')
put(ws,12,1,'双方现金流准则及capex定义仍有边界差异；PONY capex按公司Non-GAAP，未取得完整租赁附注。')

ws=wb['5b_股价与估值']
rows={2:'WRD价格USD',3:'WRD ADS等值股份（扣库存股）',4:'最新归母权益/ADS USD',5:'每ADS现金USD（非底价）',6:'每ADS流动性池USD（非底价）',7:'TTM收入人民币百万',8:'市值USD百万',9:'TTM PS',10:'H1现金池/FCF 静态月数',11:'TTM现金池/FCF 静态月数',12:'FCF消耗增50%的静态月数',13:'回购重复一次半年节奏的静态月数',14:'价格样本数',15:'低于现价的价格数值分位',16:'日样本折年',17:'ADS股本接口勾稽差',18:'研究观察价USD',19:'TTM贸易应收/收入'}
for r,name in rows.items():put(ws,r,1,name)
for r,f in {
2:'='+ref('WRD_价格'),3:'=('+ref('股份A')+'+'+ref('股份B')+'-'+ref('库存股')+')/'+ref('ADS比例'),
4:"='4_资产负债偿债'!G3*1000000/"+ref('汇率')+'/B3',5:"='4_资产负债偿债'!G4*1000000/"+ref('汇率')+'/B3',6:"='4_资产负债偿债'!G27*1000000/"+ref('汇率')+'/B3',
7:"='2_利润表五年'!H2",8:'='+ref('WRD_市值')+'/1000000',9:'=B8/(B7/'+ref('汇率')+')',
10:"='5_估值对比'!J2",11:"='4_资产负债偿债'!G27/(-'3_现金流照妖镜'!I6/"+ref('年内月数')+')',12:'=B10/'+ref('压力烧钱倍率'),
13:"='4_资产负债偿债'!G27/('3_现金流照妖镜'!B26/"+ref('半年月数')+')',
17:'=B3-'+ref('WRD_股数'),18:'='+ref('研究观察价'),19:"='4_资产负债偿债'!G9/B7"}.items():put(ws,r,2,f)
put(ws,21,1,'十年权益PS反推（2036终点；不另加现金，非目标价/非公司指引）')
for c,t in enumerate(['情景','增速下沿','增速上沿','退出PS下沿','退出PS上沿','股数累计增幅','反推价下沿USD','反推价上沿USD'],1):put(ws,22,c,t)
for r,(name,g1,g2,p1,p2,dil) in enumerate([('悲观',.10,.15,1,2,.35),('中性',.25,.30,3,4,.20),('乐观',.35,.40,5,6,.10)],23):
    put(ws,r,1,name)
    for c,(k,v) in enumerate(zip(['增速低','增速高','PS低','PS高','稀释'],[g1,g2,p1,p2,dil]),2):
        put(ws,r,c,'='+inp(name+k,v,'本报告T5探索情景假设：硬件/项目混合业务，非公司指引；PE档不适用',unit='比率/倍数',note='估算假设'))
    put(ws,r,7,f'=$B$7*(1+B{r})^{ref("期限年")}*D{r}/{ref("汇率")}*1000000/($B$3*(1+F{r}))/(1+{ref("要求回报")})^{ref("期限年")}')
    put(ws,r,8,f'=$B$7*(1+C{r})^{ref("期限年")}*E{r}/{ref("汇率")}*1000000/($B$3*(1+F{r}))/(1+{ref("要求回报")})^{ref("期限年")}')
put(ws,28,1,'现价在中性退出PS下沿/上沿的隐含十年收入CAGR')
for c,pscol in [(2,'D'),(3,'E')]:put(ws,28,c,f'=($B$2*$B$3*(1+$F$24)*(1+{ref("要求回报")})^{ref("期限年")}*{ref("汇率")}/1000000/($B$7*${pscol}$24))^(1/{ref("期限年")})-1')
put(ws,30,1,'PB档位敏感度（账面机械映射，非价值底）')
for r,pb in enumerate([.8,1,1.5,2],31):
    put(ws,r,1,'='+inp('PB档'+str(pb),pb,'本报告账面敏感性假设',unit='倍',note='假设'))
    put(ws,r,2,f'=A{r}*$B$4')
put(ws,36,1,'期权/RSU不直接加到股本：预发行激励股份可能重叠；用20%等累计稀释假设测试，不称完全摊薄。')

prices=wb.create_sheet('8_股价日线')
prices.append(['日期','WRD收盘USD（不复权）','严格低于基准价'])
with (HERE/'px_daily_WRD.csv').open() as f:
    daily=list(csv.DictReader(f))
for r,row in enumerate(daily,2):
    prices.append([row['date'],float(row['close']),f"=--(B{r}<'5b_股价与估值'!B2)"])
    prices.cell(r,2).comment=Comment(market['WRD']['daily_url']+'；2026-09-17采集','来源')
log.append(['2026-09-17','8_股价日线!B2:B'+str(prices.max_row),'逐日收盘；见CSV',market['WRD']['daily_url'],'批量逐行来源相同；不复权；每行日期为有效日'])
put(ws,14,2,f"=COUNT('8_股价日线'!B2:B{prices.max_row})")
put(ws,15,2,f"=SUM('8_股价日线'!C2:C{prices.max_row})/B14")
put(ws,16,2,'=B14/'+ref('交易日年基数'))

fmt='#,##0.000;(#,##0.000);–'
for sheet in wb:
    sheet.freeze_panes='B2'
    if sheet!=wb['0_说明']:
        sheet.column_dimensions['A'].width=38
    for c in range(2,sheet.max_column+1):sheet.column_dimensions[col(c)].width=19 if c<9 else 26
    if sheet in [raw,log]:
        sheet.column_dimensions['A'].width=30
        sheet.column_dimensions['D'].width=70
        sheet.column_dimensions['E'].width=65
    for row in sheet:
        for cell in row:
            if cell.value is None:continue
            formula=cell.data_type=='f'
            color='008000' if formula and '!' in cell.value else '000000' if formula else '0000FF' if isinstance(cell.value,(int,float)) else '25364A'
            cell.font=Font(name='Arial',size=10,color=color)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if formula or isinstance(cell.value,(int,float)):cell.number_format=fmt
            if cell.row==1:
                cell.fill=PatternFill('solid',fgColor='1F2A44');cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11)
    sheet.sheet_view.showGridLines=False
    sheet.print_options.horizontalCentered=True
for r in [12,13,14,15]:
    for c in range(2,9):wb['2_利润表五年'].cell(r,c).number_format='0.0%'
for r in [15,19,28]:
    for c in [2,3]:wb['5b_股价与估值'].cell(r,c).number_format='0.0%'
for r in range(23,26):
    for c in [2,3,6]:wb['5b_股价与估值'].cell(r,c).number_format='0.0%'
    for c in range(2,7):wb['5b_股价与估值'].cell(r,c).fill=PatternFill('solid',fgColor='FFF2CC')
for r in range(2,4):wb['5_估值对比'].cell(r,5).number_format='0.0%'
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
wb.save(OUT)

# 只支持本底稿使用的算术、比较、SUM、COUNT；拒绝未知表达式与循环引用。
cache={}; visiting=set()
ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv,ast.Pow:operator.pow}
refpattern=re.compile(r"(?:(?:'([^']+)'|([\w\u4e00-\u9fff]+))!)?(\$?[A-Z]+\$?\d+)")
def ev(sheet,address):
    address=address.replace('$',''); key=(sheet,address)
    if key in cache:return cache[key]
    if key in visiting:raise ValueError('循环引用'+str(key))
    value=wb[sheet][address].value
    if not isinstance(value,str) or not value.startswith('='):
        if not isinstance(value,(int,float,bool)):raise ValueError('非数值引用'+str(key)+str(value))
        return value
    visiting.add(key)
    expr=value[1:].replace('^','**')
    def aggr(m):
        func,body=m.group(1),m.group(2)
        a,b=body.split(':'); mt=refpattern.fullmatch(a)
        sn=mt.group(1) or mt.group(2) or sheet
        start=mt.group(3).replace('$',''); end=b.replace('$','')
        vals=[ev(sn,c.coordinate) for row in wb[sn][f'{start}:{end}'] for c in row]
        return str(sum(vals) if func=='SUM' else len(vals))
    expr=re.sub(r'(SUM|COUNT)\(([^()]+:[^()]+)\)',aggr,expr)
    expr=refpattern.sub(lambda m:'('+str(ev(m.group(1) or m.group(2) or sheet,m.group(3)))+')',expr)
    def node(n):
        if isinstance(n,ast.Constant) and isinstance(n.value,(int,float,bool)):return n.value
        if isinstance(n,ast.BinOp) and type(n.op) in ops:return ops[type(n.op)](node(n.left),node(n.right))
        if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.USub,ast.UAdd)):return -node(n.operand) if isinstance(n.op,ast.USub) else node(n.operand)
        if isinstance(n,ast.Compare) and len(n.ops)==1 and isinstance(n.ops[0],ast.Lt):return int(node(n.left)<node(n.comparators[0]))
        raise ValueError(ast.dump(n))
    result=node(ast.parse(expr,mode='eval').body)
    if not math.isfinite(result):raise ValueError('非有限值')
    cache[key]=result;visiting.remove(key);return result
for sheet in wb:
    for row in sheet:
        for cell in row:
            if cell.data_type=='f':ev(sheet.title,cell.coordinate)
checks={}
for name,addrs in {'1_分部拆解':['B9','C9','D9','E9','F9','G9'],'4_资产负债偿债':['C26','D26','E26','F26','G26'],'3_现金流照妖镜':['B27'],'5b_股价与估值':['B17']}.items():
    for a in addrs:checks[name+'!'+a]=ev(name,a)
assert all(abs(v)<(0.01 if k.endswith('B17') else 1e-4) for k,v in checks.items()),checks
assert abs(ev('5b_股价与估值','B15')-market['WRD']['price_percentile'])<1e-12
assert abs(ev('5_估值对比','G3')+162.833)<1e-9
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(OUT) as z:contents={name:z.read(name) for name in z.namelist()}
for i,sheet in enumerate(wb,1):
    name=f'xl/worksheets/sheet{i}.xml'; tree=ET.fromstring(contents[name])
    for c in tree.findall('.//s:c',ns):
        key=(sheet.title,c.attrib['r'])
        if key in cache:
            v=c.find('s:v',ns)
            if v is None:v=ET.SubElement(c,'{'+ns['s']+'}v')
            v.text=str(cache[key]);c.attrib.pop('t',None)
    contents[name]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in contents.items():z.writestr(name,data)
readback=load_workbook(OUT,data_only=True)
assert all(readback[s][a].value is not None for s,a in cache)
validation={'method':'限定公式AST逐格求值+缓存回读；非Office原生重算','formulas':len(cache),'inputs':raw.max_row-1,'errors':0,'balance_checks':checks,'anchor_checks':json.loads((HERE/'anchor_checks.json').read_text()),'snapshot_utc':datetime.fromtimestamp(market['WRD']['snapshot']['f86'],timezone.utc).isoformat(),'native_recalculation':'未执行：soffice缺失','formula_values':{s+'!'+a:v for (s,a),v in cache.items() if s!='8_股价日线'}}
(HERE/'workbook_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))
(HERE/'financial_facts.json').write_text(json.dumps({'annual':annual,'half':half,'bs':bs,'segments':segments,'pony':pony},ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in validation.items() if k not in ['formula_values','anchor_checks']},ensure_ascii=False,indent=2))
print('估值与跑道',json.dumps({a:ev('5b_股价与估值',a) for a in ['B3','B4','B5','B6','B7','B9','B10','B11','B12','B13','G23','H23','G24','H24','G25','H25','B28','C28']},ensure_ascii=False))
